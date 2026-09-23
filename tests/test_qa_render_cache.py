from __future__ import annotations

import contextlib
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import qa_gate


class QaRenderCacheTests(unittest.TestCase):
    def run_qa(self, flags: list[str], render_rc: int = 0) -> tuple[int, list[list[str]]]:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            deck = root / "deck.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(deck)
            output = root / "qa"
            calls = []

            def capture(command):
                calls.append(command)
                if Path(command[1]).name == "render_slides.py":
                    (output / "renders" / "slide-01.jpg").write_bytes(b"mock-render")
                    return render_rc, "mock render"
                return 0, ""

            argv = ["qa_gate.py", "--input", str(deck), "--outdir", str(output), *flags]
            with patch.object(sys, "argv", argv), patch.object(qa_gate, "_run", return_value=""), patch.object(qa_gate, "_run_capture", side_effect=capture), contextlib.redirect_stdout(io.StringIO()):
                return qa_gate.main(), calls

    def test_forwards_explicit_cache_directory_to_renderer(self) -> None:
        rc, calls = self.run_qa(["--render-cache-dir", "~/qa-render-cache"])
        self.assertEqual(rc, 0)
        command = next(c for c in calls if Path(c[1]).name == "render_slides.py")
        self.assertEqual(command[command.index("--cache-dir") + 1], str(Path("~/qa-render-cache").expanduser().resolve()))
        self.assertEqual(command[command.index("--dpi") + 1], "180")

    def test_cache_remains_opt_in(self) -> None:
        rc, calls = self.run_qa([])
        self.assertEqual(rc, 0)
        command = next(c for c in calls if Path(c[1]).name == "render_slides.py")
        self.assertNotIn("--cache-dir", command)
        self.assertNotIn("--cache", command)

    def test_skip_render_does_not_touch_cache(self) -> None:
        rc, calls = self.run_qa(["--render-cache-dir", "~/qa-render-cache", "--skip-render"])
        self.assertEqual(rc, 0)
        self.assertFalse(any(Path(c[1]).name == "render_slides.py" for c in calls))

    def test_render_failure_remains_blocking_with_cache_enabled(self) -> None:
        rc, _ = self.run_qa(["--render-cache-dir", "~/qa-render-cache"], render_rc=1)
        self.assertEqual(rc, 1)


if __name__ == "__main__":
    unittest.main()
