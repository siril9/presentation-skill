from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from design_rules_qa import (  # noqa: E402
    _text_role,
    check_table_caption_overlap,
    check_table_readability,
)
from finalize_quick_deck import _completion_status  # noqa: E402
from source_fidelity import check_source_fidelity  # noqa: E402
import design_rules_qa  # noqa: E402
import qa_gate  # noqa: E402


class TextRoleTests(unittest.TestCase):
    def _shape(self, *, name: str, top: float = 2.0, height: float = 0.58):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(
            Inches(1.0), Inches(top), Inches(4.0), Inches(height)
        )
        shape.name = name
        shape.text = "TABLE INDEX"
        return shape

    def test_named_metadata_is_caption_even_in_a_tall_box(self) -> None:
        shape = self._shape(name="metadata:table-index-detail")
        self.assertEqual(_text_role(shape, shape.text, 7.5), "caption")

    def test_ordinary_text_in_the_same_box_is_body(self) -> None:
        shape = self._shape(name="Text 1")
        self.assertEqual(_text_role(shape, shape.text, 7.5), "body")


class TableGeometryTests(unittest.TestCase):
    def test_row_heights_cannot_exceed_table_frame(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        table_shape = slide.shapes.add_table(
            3, 2, Inches(1.0), Inches(1.0), Inches(5.0), Inches(1.0)
        )
        for row in table_shape.table.rows:
            row.height = Inches(0.5)
        table_shape.height = Inches(1.0)

        issues = check_table_readability(0, slide, {"min_caption_pt": 7.5})

        self.assertTrue(
            any(issue["type"] == "table_rows_exceed_frame" for issue in issues)
        )

    def test_named_table_caption_cannot_overlap_table_frame(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_table(
            2, 2, Inches(1.0), Inches(1.0), Inches(5.0), Inches(1.5)
        )
        caption = slide.shapes.add_textbox(
            Inches(1.0), Inches(2.3), Inches(5.0), Inches(0.3)
        )
        caption.name = "metadata:table-caption"
        caption.text = "Source: synthetic test data"
        text_shapes = list(
            (index, shape, shape.text)
            for index, shape in enumerate(slide.shapes, start=1)
            if getattr(shape, "has_text_frame", False) and shape.text
        )

        issues = check_table_caption_overlap(0, slide, text_shapes)

        self.assertTrue(any(issue["type"] == "table_caption_overlap" for issue in issues))


class FinalizerReceiptTests(unittest.TestCase):
    def test_render_only_failure_is_deferred_without_runtime_probing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            qa_dir = Path(tmp)
            (qa_dir / "qa_report.json").write_text(
                json.dumps({"render_rc": 1, "visual_review_warning_count": 1}),
                encoding="utf-8",
            )
            status = _completion_status(
                [{"stage": "qa", "returncode": 1}], qa_dir
            )
        self.assertEqual(status["failure_category"], "render_environment")
        self.assertEqual(status["render_status"], "deferred_environment")
        self.assertIn("Do not probe", status["next_action"])

    def test_static_qa_findings_remain_source_repair_work(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            qa_dir = Path(tmp)
            (qa_dir / "qa_report.json").write_text(
                json.dumps({"render_rc": 1, "overflow_count": 1}),
                encoding="utf-8",
            )
            status = _completion_status(
                [{"stage": "qa", "returncode": 1}], qa_dir
            )
        self.assertEqual(status["failure_category"], "qa_findings")
        self.assertIn("edit outline.json", status["next_action"])


class SourceFidelityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.directory = Path(self.tmp.name).resolve()
        self.outline = self.directory / "outline.json"
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def check(self, spec):
        self.outline.write_text(json.dumps({"slides": [spec]}), encoding="utf-8")
        return check_source_fidelity(self.prs, self.outline)

    def text(self, value):
        shape = self.slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        shape.text = value
        return shape

    def chart(self, categories=("A", "B"), values=(12, 30)):
        data = CategoryChartData()
        data.categories = categories
        data.add_series("Readout", values)
        return self.slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(3), data
        ).chart

    def test_missing_lab_caption_has_original_field_pointer_and_canonical_comparison(self):
        result = self.check({"variant": "lab-run-results", "caption": "n=12; units: mg; pilot only"})
        self.assertEqual(result["error_count"], 1)
        finding = result["issues"][0]
        self.assertEqual(finding["source_pointer"], "/slides/0/caption")
        self.assertEqual(finding["slide_index"], 0)
        self.assertEqual(finding["expected_canonical"], "n=12; units: mg; pilot only")
        self.assertEqual(finding["actual_canonical"], "")
        self.assertIn("not factual truth", result["scope"])

    def test_figure_caption_slots_distinguish_sidebar_thumbnails_and_full_panels(self):
        figures = [{"caption": "Primary: n=3; synthetic only"},
                   {"caption": "Secondary: n=6; pilot only"},
                   {"caption": "Third: units mg/L"}]
        self.text(figures[0]["caption"])
        for layout in ("strip-readout", "strip_readout", "stats-strip", "metric-strip"):
            with self.subTest(layout=layout):
                self.assertEqual(self.check({"variant": "scientific-figure", "figure_layout": layout,
                                             "figures": figures})["issues"], [])
        result = self.check({"variant": "scientific-figure", "figure_layout": "panel-grid", "figures": figures})
        self.assertEqual([issue["source_pointer"] for issue in result["issues"]],
                         ["/slides/0/figures/1/caption", "/slides/0/figures/2/caption"])
        result = self.check({"variant": "image-sidebar", "figures": figures, "caption": "Visible sidebar caveat"})
        self.assertEqual([issue["source_pointer"] for issue in result["issues"]], ["/slides/0/caption"])
        self.text("Synthetic operating figure; no cu...")
        caption = "Synthetic operating figure; no customer data."
        result = self.check({"variant": "scientific-figure", "figure_layout": "ledger-rail",
                             "caption": caption, "interpretation": "No release without complete monitoring coverage.",
                             "figures": [{"caption": caption}]})
        self.assertEqual([issue["source_pointer"] for issue in result["issues"]],
                         ["/slides/0/caption", "/slides/0/interpretation", "/slides/0/figures/0/caption"])

    def test_native_figure_routes_require_primary_and_panel_captions_not_thumbnail_metadata(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("repository Node renderer unavailable")
        from PIL import Image
        image = self.directory / "panel.png"
        Image.new("RGB", (160, 100), "white").save(image)
        figures = [{"path": str(image), "caption": caption}
                   for caption in ("Primary: n=3", "Secondary: units mg/L", "Third: pilot only")]
        slides = [
            {"variant": "scientific-figure", "figure_layout": "strip-readout", "figures": figures},
            {"variant": "scientific-figure", "figure_layout": "panel-grid", "figures": figures},
            {"variant": "image-sidebar", "assets": {"hero_image": str(image)},
             "figures": figures, "caption": "Sidebar: not a power calculation"},
        ]
        self.outline.write_text(json.dumps({"slides": slides}), encoding="utf-8")
        pptx = self.directory / "figure-slots.pptx"
        result = subprocess.run([node, str(SCRIPTS / "build_deck_pptxgenjs.js"), "--outline", str(self.outline),
                                 "--output", str(pptx), "--style-preset", "lab-report"],
                                capture_output=True, text=True, timeout=30, cwd=ROOT)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        prs = Presentation(str(pptx))
        self.assertEqual(check_source_fidelity(prs, self.outline)["issues"], [])
        for index, omitted in ((0, figures[0]["caption"]), (1, figures[2]["caption"]), (2, slides[2]["caption"])):
            for shape in prs.slides[index].shapes:
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.text = run.text.replace(omitted, "")
        result = check_source_fidelity(prs, self.outline)
        self.assertEqual([issue["source_pointer"] for issue in result["issues"]],
                         ["/slides/0/figures/0/caption", "/slides/1/figures/2/caption", "/slides/2/caption"])

    def test_caption_in_speaker_notes_does_not_satisfy_visible_source(self):
        self.slide.notes_slide.notes_text_frame.text = "n=12; pilot only"
        result = self.check({"variant": "table", "caption": "n=12; pilot only"})
        self.assertEqual(result["error_count"], 1)

    def test_split_rich_runs_and_reflowed_caveat_are_retained(self):
        shape = self.text("")
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.add_run().text = "Dose 2 "
        paragraph.add_run().text = "mg"
        self.text("n=12  |  Pilot only")
        result = self.check({"variant": "table", "caption": "n=12\nPilot only",
                             "headers": [[{"text": "Dose "}, {"text": "2 "}, {"text": "mg", "options": {"bold": True}}]]})
        self.assertEqual(result["issues"], [])

    def test_changed_unit_or_denominator_is_not_a_substring_match(self):
        self.text("Dose 12 mg; n=120")
        result = self.check({"variant": "table", "caption": "2 mg", "footnotes": ["n=12"]})
        self.assertEqual(result["error_count"], 2)
        self.assertEqual(result["issues"][1]["source_pointer"], "/slides/0/footnotes/0")

    def test_decimal_units_and_denominators_are_not_substring_matches(self):
        self.text("Dose 0.2 mg; n=12.5")
        result = self.check({"variant": "table", "caption": "2 mg", "footnotes": ["n=12"]})
        self.assertEqual(result["error_count"], 2)

    def test_dropped_ordinary_slide_caveat_is_checked_in_its_body_slot(self):
        self.text("Dose was 2 mg")
        result = self.check({"variant": "standard", "bullets": ["Dose was 2 mg", "n=12; pilot only"]})
        self.assertEqual(result["error_count"], 1)
        self.assertEqual(result["issues"][0]["source_pointer"], "/slides/0/bullets/1")

    def test_timeline_caption_is_required_for_explicit_and_inferred_timelines(self):
        self.slide.notes_slide.notes_text_frame.text = "Six lots are not a power calculation."
        for variant in ({"variant": "timeline"}, {"milestones": [{"title": "Plan"}, {"title": "Review"}]}):
            with self.subTest(variant=variant):
                result = self.check({**variant, "caption": "Six lots are not a power calculation."})
                self.assertEqual(result["error_count"], 1)
                self.assertEqual(result["issues"][0]["source_pointer"], "/slides/0/caption")
        self.text("Six lots are not a power calculation.")
        self.assertEqual(self.check({"variant": "timeline", "caption": "Six lots are not a power calculation."})["issues"], [])

    def test_string_fact_labels_are_required_without_numeric_inference(self):
        for key in ("facts", "stats"):
            with self.subTest(key=key):
                result = self.check({"variant": "stats", key: ["n=3; pilot only", "72% availability"]})
                self.assertEqual(result["error_count"], 2)
                self.assertEqual(result["issues"][0]["source_pointer"], f"/slides/0/{key}/0")
                self.assertEqual(result["issues"][0]["field_role"], "fact_label")
        self.text("n=3; pilot only 72% availability")
        self.assertEqual(self.check({"variant": "stats", "facts": ["n=3; pilot only", "72% availability"]})["issues"], [])

    def test_chart_string_fact_labels_keep_inline_and_external_source_pointers(self):
        self.chart()
        chart = {"labels": ["A", "B"], "values": [12, 30],
                 "facts": ["Availability is not gardening effectiveness."]}
        result = self.check({"variant": "chart", "chart": chart})
        finding, = result["issues"]
        self.assertEqual(finding["source_pointer"], "/slides/0/chart/facts/0")
        data = self.directory / "chart-facts.json"
        data.write_text(json.dumps(chart), encoding="utf-8")
        result = self.check({"variant": "chart", "assets": {"chart_data": data.name}})
        finding, = result["issues"]
        self.assertEqual(finding["source_pointer"], "/slides/0/assets/chart_data")
        self.assertEqual(finding["source_data_file"], str(data))
        self.assertEqual(finding["source_data_pointer"], "/facts/0")
        # An explicit slide-level list replaces payload facts, even if empty.
        self.assertEqual(self.check({"variant": "chart", "chart": chart, "facts": []})["issues"], [])
        self.text(chart["facts"][0])
        self.assertEqual(self.check({"variant": "chart", "chart": chart})["issues"], [])

    def test_optional_notes_metadata_and_reference_headers_do_not_fail(self):
        result = self.check({"role": "references", "variant": "table",
                             "headers": ["ID", "Source"], "rows": [["D1", "A long reformatted source"]],
                             "notes": "Presenter-only caveat", "sources": [{"id": "source-id"}],
                             "style": {"caption": "not user content"}})
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["checked_count"], 0)

    def test_references_treatment_allows_bullets_reformatted_as_a_register(self):
        self.text("S1 Synthetic incident notes")
        result = self.check({"variant": "standard", "treatment_key": "references",
                             "bullets": ["S1: Synthetic incident notes"], "body": "Optional references explanation"})
        self.assertEqual(result["issues"], [])

    def test_policy_reference_footnote_metadata_is_not_an_ordinary_visible_caveat(self):
        footnotes = ["SOURCE REGISTER", "S1 | Neighborhood evidence", "S2 | Decision evidence"]
        spec = {"role": "references", "variant": "table", "table": {"footnotes": footnotes}}
        self.assertEqual(self.check(spec)["error_count"], 3)
        marker = self.text("ACCOUNTABILITY REGISTER")
        marker.name = "metadata:policy-accountability-register"
        self.assertEqual(self.check(spec)["issues"], [])
        ordinary = {"role": "table", "variant": "table", "table": {"footnotes": ["n=12; pilot only"]}}
        result = self.check(ordinary)
        self.assertEqual(result["error_count"], 1)
        self.assertEqual(result["issues"][0]["source_pointer"], "/slides/0/table/footnotes/0")
        with_caption = {**spec, "table": {"footnotes": footnotes, "caption": "Critical caption remains required"}}
        result = self.check(with_caption)
        self.assertEqual(result["error_count"], 1)
        self.assertEqual(result["issues"][0]["source_pointer"], "/slides/0/table/caption")

    def test_native_js_inferred_chart_and_hidden_axis_titles_match_checker(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("repository Node renderer unavailable")
        pptx = self.directory / "native.pptx"
        self.outline.write_text(json.dumps({"slides": [
            {"type": "content", "visual_intent": "data", "title": "Results", "chart": {
                "labels": ["A", "B"], "values": [12, 30],
                "facts": ["Pilot availability measure; not a measure of gardening outcomes."], "options": {
                    "catAxisTitle": "Cohort denominator", "valAxisTitle": "mg/L",
                    "showCatAxisTitle": False, "showValAxisTitle": False}}},
            {"type": "title", "variant": "chart", "title": "Summary",
             "chart": {"labels": ["A", "B"], "values": [12, 30]}}
        ]}), encoding="utf-8")
        result = subprocess.run([node, str(SCRIPTS / "build_deck_pptxgenjs.js"), "--outline", str(self.outline),
                                 "--output", str(pptx), "--style-preset", "lab-report"],
                                capture_output=True, text=True, timeout=30, cwd=ROOT)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        fidelity = check_source_fidelity(Presentation(str(pptx)), self.outline)
        self.assertEqual(fidelity["issues"], [])

    def test_chart_values_are_compared_numerically_not_as_formatted_text(self):
        self.chart(values=(1000, 30))
        result = self.check({"variant": "chart", "chart": {
            "categories": ["A", "B"], "series": [{"values": ["1,000", "30%"]}],
            "options": {"dataLabelFormatCode": "0.0%"}}})
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["checked_count"], 3)

    def test_native_axis_title_defaults_follow_v2_not_archival_or_axisless_routes(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("repository Node renderer unavailable")
        slides = [{"variant": "chart", "title": "Results", "chart": {
            "type": kind, "labels": ["A", "B"], "values": [12, 30],
            "options": {"catAxisTitle": "Cohort denominator", "valAxisTitle": "mg/L"}}}
            for kind in ("bar", "pie", "doughnut")]
        slides.append({**slides[0], "role": "evidence"})
        pin = {"renderer_role_systems_v1": {"data_system_id": "data-assay-readout"}}
        for config in ({}, {"metadata": pin}, {"deck_style": pin}):
            with self.subTest(config=config):
                pptx = self.directory / "native-axis.pptx"
                self.outline.write_text(json.dumps({**config, "slides": slides}), encoding="utf-8")
                result = subprocess.run([node, str(SCRIPTS / "build_deck_pptxgenjs.js"),
                                         "--outline", str(self.outline), "--output", str(pptx),
                                         "--style-preset", "lab-report"],
                                        capture_output=True, text=True, timeout=30, cwd=ROOT)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                prs = Presentation(str(pptx))
                self.assertEqual(check_source_fidelity(prs, self.outline)["issues"], [])
                for index, slide in enumerate(prs.slides):
                    chart = next(shape.chart._chartSpace for shape in slide.shapes if shape.has_chart)
                    titles = chart.xpath(".//c:catAx/c:title | .//c:valAx/c:title")
                    self.assertEqual(len(titles), 2 if not config and index == 0 else 0)
                    for title in titles:
                        title.getparent().remove(title)
                findings = check_source_fidelity(prs, self.outline)["issues"]
                self.assertEqual([issue["source_pointer"] for issue in findings],
                                 ["/slides/0/chart/options/catAxisTitle", "/slides/0/chart/options/valAxisTitle"]
                                 if not config else [])

    def test_deliberately_hidden_axis_titles_are_not_required(self):
        self.chart()
        result = self.check({"variant": "chart", "chart": {
            "categories": ["A", "B"], "values": [12, 30],
            "options": {"catAxisTitle": "Cohort denominator", "valAxisTitle": "mg/L",
                        "showCatAxisTitle": False, "showValAxisTitle": False}}})
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["checked_count"], 3)

    def test_visible_axis_titles_still_block_when_lost(self):
        self.chart()
        for flags in ({}, {"showCatAxisTitle": True, "showValAxisTitle": True}):
            with self.subTest(flags=flags):
                result = self.check({"variant": "chart", "chart": {
                    "categories": ["A", "B"], "values": [12, 30],
                    "options": {"catAxisTitle": "Cohort denominator", "valAxisTitle": "mg/L", **flags}}})
                self.assertEqual(result["error_count"], 2)
                self.assertEqual(result["issues"][1]["source_pointer"], "/slides/0/chart/options/valAxisTitle")

    def test_inferred_asset_chart_uses_native_data_and_ignores_unused_body(self):
        self.chart()
        data = self.directory / "chart.json"
        data.write_text(json.dumps({"labels": ["A", "B"], "values": [12, 30]}), encoding="utf-8")
        for slide_type in ("content", "text"):
            with self.subTest(type=slide_type):
                result = self.check({"type": slide_type, "visual_intent": "data", "assets": {"chart_data": "chart.json"},
                                     "body": "unused by chart layout"})
                self.assertEqual(result["issues"], [])
                self.assertEqual(result["checked_count"], 3)

    def test_title_type_overrides_chart_variant_without_requiring_chart_data(self):
        self.text("Results")
        result = self.check({"type": "title", "variant": "chart", "title": "Results", "body": "optional title metadata",
                             "chart": {"labels": ["A", "B"], "values": [12, 30]}})
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["checked_count"], 1)

    def test_inferred_table_caption_still_blocks(self):
        self.text("Series Value A 12")
        result = self.check({"type": "content", "headers": ["Series", "Value"], "rows": [["A", 12]],
                             "caption": "n=12; pilot only", "body": "unused table prose"})
        self.assertEqual(result["error_count"], 1)
        self.assertEqual(result["issues"][0]["source_pointer"], "/slides/0/caption")

    def test_payload_does_not_imply_chart_without_data_intent(self):
        self.text("Caveat: pilot only")
        result = self.check({"body": "Caveat: pilot only", "chart": {"labels": ["A", "B"], "values": [12, 30]}})
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["checked_count"], 1)

    def test_resolved_workspace_keeps_source_asset_root_and_original_pointers(self):
        self.chart()
        (self.directory / "workspace.json").write_text(json.dumps({"outline": "outline.json"}), encoding="utf-8")
        (self.directory / "asset_plan.json").write_text(json.dumps({"charts": [{"name": "run", "path": "chart.json"}]}), encoding="utf-8")
        (self.directory / "chart.json").write_text(json.dumps({"labels": ["A", "B"], "values": [12, 31]}), encoding="utf-8")
        self.check({"assets": {"chart_data": "chart:run"}, "body": "unused after resolution"})
        resolved = self.directory / "build/outline_resolved.json"
        resolved.parent.mkdir()
        resolved.write_text(json.dumps({"slides": [{"variant": "chart", "assets": {"chart_data": "chart:run"}}]}), encoding="utf-8")
        result = check_source_fidelity(self.prs, resolved)
        self.assertEqual(result["error_count"], 1)
        finding = result["issues"][0]
        self.assertEqual(finding["source_file"], str(self.outline))
        self.assertEqual(finding["source_pointer"], "/slides/0/assets/chart_data")
        self.assertEqual(finding["source_data_pointer"], "/values")
        alternate = self.directory / "alternate-root"
        alternate.mkdir()
        (alternate / "asset_plan.json").write_text(json.dumps({"charts": [{"name": "run", "path": "chart.json"}]}), encoding="utf-8")
        data = alternate / "chart.json"
        data.write_text(json.dumps({"labels": ["A", "B"], "values": [12, 30]}), encoding="utf-8")
        self.assertEqual(check_source_fidelity(self.prs, resolved, asset_root=alternate)["issues"], [])
        data.write_text(json.dumps({"labels": ["A", "B"], "values": [12, 32]}), encoding="utf-8")
        finding, = check_source_fidelity(self.prs, resolved, asset_root=alternate)["issues"]
        self.assertEqual(finding["source_file"], str(self.outline))
        self.assertEqual(finding["source_pointer"], "/slides/0/assets/chart_data")
        self.assertEqual(finding["source_data_file"], str(data))
        self.assertEqual(finding["source_data_pointer"], "/values")

    def test_asset_root_alias_and_relative_sources_flow_through_existing_gate(self):
        self.chart()
        self.text("Sample Signal A 12 n=12")
        root = self.directory / "assets-root"
        root.mkdir()
        data = root / "data.json"
        self.outline = self.directory / "audition/outline.json"
        self.outline.parent.mkdir()
        (root / "asset_plan.json").write_text(json.dumps({
            "charts": [{"name": "run", "path": "data.json"}],
            "tables": [{"name": "run", "path": "data.json"}]}), encoding="utf-8")
        pptx = self.directory / "deck.pptx"
        self.prs.save(pptx)
        for variant, payload, field in (
            ("chart", {"labels": ["A", "B"], "values": [12, 30]}, "/values"),
            ("table", {"headers": ["Sample", "Signal"], "rows": [["A", 12]], "caption": "n=12"}, "/caption"),
        ):
            for reference in ("data.json", variant + ":run"):
                with self.subTest(variant=variant, reference=reference):
                    data.write_text(json.dumps(payload), encoding="utf-8")
                    spec = {"type": "content", "visual_intent": "data", variant: reference}
                    self.check(spec)
                    self.assertEqual(check_source_fidelity(self.prs, self.outline, asset_root=root)["issues"], [])
                    calls = []

                    def run(cmd):
                        calls.append(Path(cmd[1]).name)
                        if Path(cmd[1]).name == "design_rules_qa.py":
                            self.assertEqual(cmd[cmd.index("--asset-root") + 1], str(root))
                            with patch.object(sys, "argv", cmd[1:]), patch("builtins.print"):
                                return design_rules_qa.main(), "design report written"
                        return 0, "[]" if Path(cmd[1]).name == "visual_qa.py" else ""

                    output = self.directory / "qa"
                    argv = ["qa_gate.py", "--input", str(pptx), "--outline", str(self.outline),
                            "--asset-root", str(root), "--outdir", str(output), "--skip-render"]
                    with patch.object(sys, "argv", argv), patch.object(qa_gate, "_run_capture", side_effect=run), patch("builtins.print"):
                        self.assertEqual(qa_gate.main(), 0)
                    self.assertEqual(calls.count("design_rules_qa.py"), 1)
                    report = json.loads((output / "qa_report.json").read_text(encoding="utf-8"))
                    self.assertEqual(report["source_fidelity_error_count"], 0)
                    changed = {**payload, **({"values": [12, 31]} if variant == "chart" else {"caption": "n=13"})}
                    data.write_text(json.dumps(changed), encoding="utf-8")
                    finding, = check_source_fidelity(self.prs, self.outline, asset_root=root)["issues"]
                    self.assertEqual(finding["source_file"], str(self.outline))
                    self.assertEqual(finding["source_pointer"], "/slides/0/" + variant)
                    self.assertEqual(finding["source_data_file"], str(data))
                    self.assertEqual(finding["source_data_pointer"], field)

    def test_flat_chart_and_per_series_labels_are_supported(self):
        self.chart()
        for payload in ({"labels": ["A", "B"], "values": [12, 30]},
                        {"series": [{"labels": ["A", "B"], "values": [12, 30]}]}):
            with self.subTest(payload=payload):
                self.assertEqual(self.check({"variant": "chart", "chart": payload})["issues"], [])

    def test_swapped_chart_values_and_categories_are_both_errors(self):
        self.chart(categories=("B", "A"), values=(30, 12))
        result = self.check({"variant": "chart", "chart": {
            "categories": ["A", "B"], "series": [{"values": [12, 30]}]}})
        self.assertEqual(result["error_count"], 2)
        values, labels = result["issues"]
        self.assertEqual(values["source_pointer"], "/slides/0/chart/series/0/values")
        self.assertEqual(values["expected_canonical"], [12.0, 30.0])
        self.assertEqual(values["actual_canonical"], [30.0, 12.0])
        self.assertEqual(labels["source_pointer"], "/slides/0/chart/categories")

    def test_chart_reference_keeps_original_pointer_and_external_data_pointer(self):
        data = self.directory / "chart.json"
        data.write_text(json.dumps({"labels": ["A", "B"], "values": [12, 31]}), encoding="utf-8")
        self.chart()
        result = self.check({"variant": "chart", "assets": {"chart_data": "chart.json"}})
        finding = result["issues"][0]
        self.assertEqual(finding["source_pointer"], "/slides/0/assets/chart_data")
        self.assertEqual(finding["source_file"], str(self.outline))
        self.assertEqual(finding["source_data_file"], str(data))
        self.assertEqual(finding["source_data_pointer"], "/values")

    def test_table_reference_caption_and_inline_override_have_correct_provenance(self):
        data = self.directory / "table.json"
        data.write_text(json.dumps({"caption": "old caption", "footnotes": ["pilot only"]}), encoding="utf-8")
        result = self.check({"variant": "table", "table": "table.json", "caption": "new caption"})
        caption, footnote = result["issues"]
        self.assertEqual(caption["source_pointer"], "/slides/0/caption")
        self.assertEqual(caption["source_data_file"], str(self.outline))
        self.assertEqual(footnote["source_pointer"], "/slides/0/table")
        self.assertEqual(footnote["source_data_file"], str(data))
        self.assertEqual(footnote["source_data_pointer"], "/footnotes/0")

    def test_missing_native_chart_is_blocking_even_if_values_appear_in_text(self):
        self.text("A B 12 30")
        result = self.check({"variant": "chart", "chart": {"labels": ["A", "B"], "values": [12, 30]}})
        self.assertEqual(result["error_count"], 3)
        self.assertEqual(result["issues"][0]["type"], "source_fidelity_chart_series_mismatch")

    def test_chart_readout_fallback_does_not_require_overridden_optional_notes(self):
        self.chart()
        self.text("Pilot only")
        result = self.check({"variant": "chart", "interpretation": "Pilot only", "notes": "speaker only",
                             "chart": {"labels": ["A", "B"], "values": [12, 30], "notes": "overridden note"}})
        self.assertEqual(result["issues"], [])

    def test_invalid_outline_is_an_explicit_audit_error_and_legacy_no_source_is_disabled(self):
        self.outline.write_text("{", encoding="utf-8")
        result = check_source_fidelity(self.prs, self.outline)
        self.assertEqual(result["issues"][0]["type"], "source_fidelity_audit_failed")
        result = check_source_fidelity(self.prs, None)
        self.assertFalse(result["enabled"])
        self.assertEqual(result["issues"], [])

    def test_existing_design_operation_gates_missing_caption(self):
        pptx = self.directory / "deck.pptx"
        report = self.directory / "design_rules.json"
        self.prs.save(pptx)
        self.check({"variant": "lab-run-results", "caption": "n=12; pilot only"})
        argv = ["design_rules_qa.py", "--input", str(pptx), "--outline", str(self.outline), "--report", str(report)]
        with patch.object(sys, "argv", argv), patch("builtins.print"):
            self.assertEqual(design_rules_qa.main(), 1)
        payload = json.loads(report.read_text(encoding="utf-8"))
        self.assertEqual(payload["source_fidelity_error_count"], 1)
        self.assertEqual(payload["issues"][0]["source_pointer"], "/slides/0/caption")

    def test_normal_gate_defers_office_but_keeps_reports_and_nonzero_exit(self):
        pptx = self.directory / "deck.pptx"
        output = self.directory / "qa"
        self.prs.save(pptx)
        self.check({"variant": "lab-run-results", "caption": "n=12; pilot only"})
        calls = []

        def run(cmd):
            calls.append(Path(cmd[1]).name)
            if Path(cmd[1]).name == "design_rules_qa.py":
                with patch.object(sys, "argv", cmd[1:]), patch("builtins.print"):
                    return design_rules_qa.main(), "design report written"
            if Path(cmd[1]).name == "visual_qa.py":
                return 0, "[]"
            return 0, ""

        argv = ["qa_gate.py", "--input", str(pptx), "--outline", str(self.outline),
                "--outdir", str(output), "--run-visual-review"]
        with patch.object(sys, "argv", argv), patch.object(qa_gate, "_run_capture", side_effect=run), patch("builtins.print"):
            self.assertEqual(qa_gate.main(), 1)
        payload = json.loads((output / "qa_report.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["source_fidelity_error_count"], 1)
        self.assertEqual(payload["render_status"], "deferred_source_fidelity")
        self.assertEqual(payload["render_deferred_reason"], "source_fidelity_findings")
        self.assertIsNone(payload["render_rc"])
        self.assertNotIn("render_slides.py", calls)
        self.assertNotIn("visual_review.py", calls)
        self.assertIn("design_rules_qa.py", calls)
        self.assertTrue((output / "design_rules.json").exists())
        self.assertTrue((output / "visual_review/visual_review.json").exists())

    def test_real_gate_cli_blocks_missing_caption_before_office(self):
        pptx = self.directory / "deck.pptx"
        output = self.directory / "qa"
        self.text("Results")
        self.prs.save(pptx)
        self.check({"variant": "lab-run-results", "title": "Results", "caption": "n=12; pilot only"})
        command = [sys.executable, str(SCRIPTS / "qa_gate.py"), "--input", str(pptx),
                   "--outline", str(self.outline), "--outdir", str(output), "--run-visual-review"]
        result = subprocess.run(command, capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        report = json.loads((output / "qa_report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["source_fidelity_error_count"], 1)
        self.assertIsNone(report["render_rc"])
        self.assertEqual(report["rendered_slide_count"], 0)
        self.assertEqual(report["render_deferred_reason"], "source_fidelity_findings")


if __name__ == "__main__":
    unittest.main()
