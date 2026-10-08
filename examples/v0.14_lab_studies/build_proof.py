#!/usr/bin/env python3
"""Publish compact evidence from reviewed renders, not invented slide mockups."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
sys.path.insert(0, str(REPO / "scripts"))
from office_package_hash import office_package_normalized_sha256
from visual_review_receipt import validate_receipt

STUDIES = ("white_calibration", "dark_contrast", "light_river_report")
SELECTED = (("white_calibration", 4), ("dark_contrast", 2), ("light_river_report", 4),
            ("white_calibration", 3), ("dark_contrast", 5), ("light_river_report", 3))
REFERENCES = (
    ("lab-residual-panels", "white_calibration", 4, "scientific-figure",
     ["lab", "residual", "paired", "calibration", "figure"],
     "Two comparable residual plots, identical scales, full panel captions and a separate interpretation.",
     ["scientific-evidence-plate", "clinical-care-pathway", "technical-telemetry-canvas"]),
    ("lab-native-method", "white_calibration", 3, "flow",
     ["lab", "method", "workflow", "process", "holdout"],
     "Four ordered native stages with equal rectangles, editable connectors and a separate provenance caption.",
     ["scientific-evidence-plate", "operations-grid", "technical-telemetry-canvas"]),
)


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def contact(paths: list[Path], output: Path, columns: int, width: int) -> None:
    gutter = 20
    tile_w = (width - gutter * (columns + 1)) // columns
    tile_h = round(tile_w * 9 / 16)
    rows = (len(paths) + columns - 1) // columns
    sheet = Image.new("RGB", (width, rows * tile_h + (rows + 1) * gutter), "#E6E8EA")
    for index, path in enumerate(paths):
        with Image.open(path) as source:
            tile = source.convert("RGB")
            tile.thumbnail((tile_w, tile_h), Image.Resampling.LANCZOS)
            x = gutter + (index % columns) * (tile_w + gutter) + (tile_w - tile.width) // 2
            y = gutter + (index // columns) * (tile_h + gutter) + (tile_h - tile.height) // 2
            sheet.paste(tile, (x, y))
    sheet.save(output, quality=91, optimize=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--build-root", type=Path, required=True)
    parser.add_argument("--zip", type=Path, required=True)
    args = parser.parse_args()
    build = args.build_root.resolve()
    records = []
    for study in STUDIES:
        folder = build / study
        pptx = folder / "deck.pptx"
        qa = folder / "qa"
        receipt = json.loads((qa / "finalize_receipt.json").read_text())
        if not receipt["passed"] or any(receipt["qa_counts"].values()):
            raise ValueError(f"{study}: automated gates are not clean")
        verdict = validate_receipt(receipt_path=qa / "visual_review_receipt.json",
                                   pptx_path=pptx, renders_dir=qa / "renders", fail_on_warnings=True)
        if not verdict["passed"]:
            raise ValueError(f"{study}: exact visual-review receipt is not valid: {verdict}")
        review_copy = ROOT / "reviews" / f"{study}.json"
        review_copy.parent.mkdir(exist_ok=True)
        shutil.copyfile(qa / "visual_review_receipt.json", review_copy)
        outline = json.loads((ROOT / study / "outline.json").read_text())
        prs = Presentation(pptx)
        renders = [qa / "renders" / f"slide-{index:02d}.jpg" for index in range(1, len(prs.slides) + 1)]
        contact(renders, ROOT / f"{study}_contact_sheet.jpg", 2, 2000)
        records.append({
            "study": study, "title": outline["title"], "source": f"{study}/outline.json",
            "source_sha256": sha(ROOT / study / "outline.json"), "slide_count": len(prs.slides),
            "pptx_sha256": sha(pptx), "normalized_pptx_sha256": office_package_normalized_sha256(pptx),
            "automated_qa_counts": receipt["qa_counts"], "visual_review": "pass",
            "native_charts": sum(shape.has_chart for slide in prs.slides for shape in slide.shapes),
            "native_tables": sum(shape.has_table for slide in prs.slides for shape in slide.shapes),
            "native_method_slides": [index + 1 for index, slide in enumerate(outline["slides"]) if slide.get("flow_steps")],
            "figure_layouts": sorted({slide["figure_layout"] for slide in outline["slides"] if "figure_layout" in slide}),
            "render_sha256": [sha(path) for path in renders],
            "visual_review_receipt_sha256": sha(qa / "visual_review_receipt.json"),
            "visual_review_receipt": f"reviews/{study}.json",
            "pipeline_seconds": receipt["total_duration_seconds"],
            "timing_scope": receipt["timing_scope"],
        })
    contact([build / study / "qa/renders" / f"slide-{index:02d}.jpg" for study, index in SELECTED],
            REPO / "examples/v0.14_lab_showcase.jpg", 3, 2400)
    manifest = {
        "schema_version": "lab_release_proof/v1", "release": "v0.14.0", "synthetic": True,
        "scope": "Three original design studies, not a model or cross-generator benchmark. Figures are images; text, charts, tables and native methods remain editable.",
        "core_source_sha256": {name: sha(REPO / name) for name in (
            "templates/pptxgenjs/slides.js", "scripts/build_deck_pptxgenjs.js",
            "scripts/preflight.py", "scripts/source_fidelity.py", "scripts/design_rules_qa.py", "scripts/inventory.py",
            "references/renderer_role_contracts_v2.json")},
        "selection": [{"study": study, "slide": index} for study, index in SELECTED],
        "studies": records,
    }
    (ROOT / "release_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    catalog_path = REPO / "references/visual_reference_catalog.json"
    catalog = json.loads(catalog_path.read_text())
    for reference_id, study, number, variant, keywords, composition, grammars in REFERENCES:
        source = build / study / "qa/renders" / f"slide-{number:02d}.jpg"
        relative = f"references/assets/visual_references/{reference_id}.jpg"
        destination = REPO / relative
        with Image.open(source) as original:
            snapshot = original.convert("RGB")
            snapshot.thumbnail((1200, 675), Image.Resampling.LANCZOS)
            canvas = Image.new("RGB", (1200, 675), snapshot.getpixel((0, 0)))
            canvas.paste(snapshot, ((1200 - snapshot.width) // 2, (675 - snapshot.height) // 2))
            canvas.save(destination, quality=70, subsampling=0, optimize=True)
            dimensions = canvas.size
        outline = json.loads((ROOT / study / "outline.json").read_text())
        record = {
            "id": reference_id, "roles": ["evidence"], "content_shapes": [variant],
            "purposes": keywords, "keywords": keywords, "density": "medium",
            "composition": composition,
            "mixing": "Borrow the evidence allocation, not the synthetic facts or a fixed slide sequence.",
            "compatible_grammars": grammars, "path": relative,
            "source_grammar": outline["deck_style"]["composition_grammar"],
            "snapshot": {"bytes": destination.stat().st_size, "sha256": sha(destination),
                         "width": dimensions[0], "height": dimensions[1]},
            "provenance": {
                "study": study, "slide": number, "source_outline": f"examples/v0.14_lab_studies/{study}/outline.json",
                "source_outline_sha256": sha(ROOT / study / "outline.json"),
                "source_render": f"v0.14.0 release: {study}.pptx slide {number}",
                "source_render_sha256": sha(source), "source_role": "evidence", "source_variant": variant,
                "source_title": outline["slides"][number - 1]["title"],
                "review": "examples/v0.14_lab_studies/release_manifest.json", "finalize_passed": True,
            },
        }
        catalog["exemplars"] = [item for item in catalog["exemplars"] if item["id"] != reference_id] + [record]
    catalog["origin"] = "Original repository Sol design studies: September 2026 and GPT-6.1 Sol lab studies, October 2026. Synthetic scenarios, not measured findings or a benchmark."
    catalog["asset_bytes"] = sum(item["snapshot"]["bytes"] for item in catalog["exemplars"])
    catalog_path.write_text(json.dumps(catalog, indent=2) + "\n")
    sources = [path for path in ROOT.rglob("*") if path.is_file()
               and not any(part.startswith(".") or part in {"build", "qa", "__pycache__"}
                           for part in path.relative_to(ROOT).parts)]
    args.zip.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(args.zip, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        entries = {f"source/{path.relative_to(ROOT).as_posix()}": path for path in sources}
        entries.update({f"{study}.pptx": build / study / "deck.pptx" for study in STUDIES})
        entries["showcase.jpg"] = REPO / "examples/v0.14_lab_showcase.jpg"
        for name, path in sorted(entries.items()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.external_attr = 0o100644 << 16
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, path.read_bytes())
    print(json.dumps({"studies": len(records), "slides": sum(item["slide_count"] for item in records),
                      "zip": str(args.zip), "sha256": sha(args.zip)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
