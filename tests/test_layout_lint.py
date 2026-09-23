from __future__ import annotations

import sys
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from layout_lint import ShapeInfo, _rectangle_union_area, _has_display_section_title  # noqa: E402
from pptx import Presentation
from pptx.util import Inches, Pt


def _shape(left: float, top: float, width: float, height: float) -> ShapeInfo:
    return ShapeInfo(
        shape_id="test",
        name="test",
        left=left,
        top=top,
        width=width,
        height=height,
        has_text=False,
        is_auto_shape=True,
        auto_shape_type_name="RECTANGLE",
        is_line_like=False,
        area=width * height,
    )


class RectangleUnionAreaTests(unittest.TestCase):
    def test_section_whitespace_needs_actual_display_title(self) -> None:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        outline = {"role": "section", "title": "Evidence and next steps"}
        self.assertFalse(_has_display_section_title(slide, outline))
        shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = outline["title"]
        run.font.size = Pt(12)
        self.assertFalse(_has_display_section_title(slide, outline))
        run.font.size = Pt(36)
        self.assertTrue(_has_display_section_title(slide, outline))
        self.assertFalse(_has_display_section_title(slide, {"role": "section", "title": "Missing title"}))
        self.assertFalse(_has_display_section_title(slide, {"title": outline["title"]}))

    def test_nested_shape_is_not_double_counted(self) -> None:
        shapes = [_shape(1, 1, 4, 3), _shape(2, 2, 1, 1)]
        self.assertAlmostEqual(
            _rectangle_union_area(shapes, slide_w=10, slide_h=7.5),
            12.0,
        )

    def test_overlapping_shapes_count_only_visible_union(self) -> None:
        shapes = [_shape(0, 0, 4, 2), _shape(2, 0, 4, 2)]
        self.assertAlmostEqual(
            _rectangle_union_area(shapes, slide_w=10, slide_h=7.5),
            12.0,
        )

    def test_shapes_are_clipped_to_slide_bounds(self) -> None:
        shapes = [_shape(-1, -1, 3, 3), _shape(9, 7, 3, 2)]
        self.assertAlmostEqual(
            _rectangle_union_area(shapes, slide_w=10, slide_h=7.5),
            4.5,
        )


if __name__ == "__main__":
    unittest.main()
