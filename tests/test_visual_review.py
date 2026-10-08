from __future__ import annotations

import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from visual_review import _analyze_text_shapes, _measurement_font, _shape_record  # noqa: E402
from inventory import _overflow_amount  # noqa: E402


def _text(
    slide,
    text: str,
    y: float,
    h: float,
    font_pt: float,
    *,
    x: float = 0.5,
    w: float = 9.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(font_pt)
    return shape


class VisualReviewTests(unittest.TestCase):
    def _measured_shape(self, text: str, h: float, *, w: float = 4.2):
        face = next((face for face in ("Helvetica Neue", "Arial", "DejaVu Sans", "Liberation Sans")
                     if _measurement_font(face, False, False, 128) is not None
                     and _measurement_font(face, True, False, 128) is not None), None)
        if face is None:
            self.skipTest("No declared test font is installed for Pillow measurement")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = _text(slide, text, 2.0, h, 16, w=w)
        shape.text_frame.margin_left = shape.text_frame.margin_right = 0
        shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = face
        return presentation, shape, face

    def test_declared_font_preserves_fitting_explicit_lines_and_bold_prefix(self) -> None:
        presentation, shape, face = self._measured_shape("", 1.29)
        paragraph = shape.text_frame.paragraphs[0]
        prefix = paragraph.add_run()
        prefix.text = "Release / commissioning:"
        prefix.font.name, prefix.font.size, prefix.font.bold = face, Pt(16), True
        body = paragraph.add_run()
        body.text = " Require\ndocumented site tests and authorized\nsign-off. No sign-off means no service\nrelease."
        body.font.name, body.font.size = face, Pt(16)
        # The original empty run has no visible text and should not require a font.
        record = _shape_record(1, 1, shape)
        self.assertEqual(len(record["estimated_lines"]), 4)
        self.assertAlmostEqual(record["estimated_height"], 4 * 16 / 72 * 1.28)
        self.assertEqual(_overflow_amount(shape, shape.text), 0)
        self.assertNotIn("text_box_clip_risk", {item["type"] for item in _analyze_text_shapes(presentation)})
        shape.height = Inches(0.65)
        self.assertGreater(_overflow_amount(shape, shape.text), 0)
        self.assertIn("text_box_clip_risk", {item["type"] for item in _analyze_text_shapes(presentation)})

    def test_wide_unbroken_text_still_warns_with_real_font(self) -> None:
        presentation, _, _ = self._measured_shape("W" * 50, 0.5, w=2.0)
        self.assertIn("text_box_clip_risk", {item["type"] for item in _analyze_text_shapes(presentation)})

    def test_mixed_run_sizes_use_large_body_height_not_first_metadata_size(self) -> None:
        presentation, shape, face = self._measured_shape("Metadata\n", 0.65)
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.runs[0].font.size = Pt(9)
        body = paragraph.add_run()
        body.text = "Full-size body\nSecond body line\nThird body line"
        body.font.name, body.font.size = face, Pt(16)
        self.assertIn("text_box_clip_risk", {item["type"] for item in _analyze_text_shapes(presentation)})

    def test_native_margins_and_paragraph_spacing_count_toward_clip_height(self) -> None:
        presentation, shape, _ = self._measured_shape("Single line", 0.4)
        shape.text_frame.margin_top = shape.text_frame.margin_bottom = Inches(0.1)
        shape.text_frame.paragraphs[0].space_after = Pt(12)
        self.assertIn("text_box_clip_risk", {item["type"] for item in _analyze_text_shapes(presentation)})

    def test_unavailable_font_keeps_conservative_clip_check(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = _text(slide, "First line\nSecond line\nThird line\nFourth line", 2.0, 0.5, 16)
        shape.text_frame.paragraphs[0].runs[0].font.name = "Not installed QA font"
        with patch("visual_review._measurement_font", return_value=None):
            self.assertGreater(_overflow_amount(shape, shape.text), 0)
            self.assertIn("text_box_clip_risk", {item["type"] for item in _analyze_text_shapes(presentation)})

    def test_unproblematic_body_box_does_not_load_font_metrics(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _text(slide, "Short ordinary text", 2.0, 1.0, 16)
        with patch("visual_review._measured_text_layout") as measure:
            _analyze_text_shapes(presentation)
            shape = presentation.slides[0].shapes[0]
            self.assertEqual(_overflow_amount(shape, shape.text), 0)
            measure.assert_not_called()

    def test_unreadable_metrics_fall_back_without_losing_clip_warning(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _text(slide, "First line\nSecond line\nThird line", 2.0, 0.3, 16)
        with patch("visual_review._measured_text_layout", side_effect=OSError("unreadable font")):
            self.assertIn("text_box_clip_risk", {item["type"] for item in _analyze_text_shapes(presentation)})

    def test_footer_clearance_does_not_compare_a_low_register_to_itself(self) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(10.0)
        presentation.slide_height = Inches(5.625)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _text(slide, "Assay readout", 0.2, 0.5, 28)
        _text(slide, "Source: data/assay.csv", 4.405, 0.30, 9)
        _text(slide, "METHOD | CONTROL | n | UNCERTAINTY", 4.865, 0.16, 8)
        _text(slide, "Sources: data/assay.csv", 5.325, 0.30, 8)

        issues = _analyze_text_shapes(presentation)

        self.assertNotIn("footer_clearance_risk", {item["type"] for item in issues})

    def test_shrink_to_fit_pressure_is_information_not_false_clip_warning(self) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(10.0)
        presentation.slide_height = Inches(5.625)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = _text(
            slide,
            "A long editable sentence that intentionally relies on PowerPoint shrink to fit.",
            2.0,
            0.30,
            12,
            w=3.0,
        )
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

        issues = _analyze_text_shapes(presentation)

        pressure = [item for item in issues if item["type"] == "text_autofit_pressure"]
        self.assertTrue(pressure)
        self.assertTrue(all(item["severity"] == "info" for item in pressure))
        self.assertNotIn("text_box_clip_risk", {item["type"] for item in issues})

    def test_title_clearance_ignores_non_overlapping_side_rail_text(self) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(10.0)
        presentation.slide_height = Inches(5.625)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _text(slide, "Main title", 0.3, 0.7, 28, x=0.5, w=5.0)
        _text(slide, "SIDE RAIL", 0.5, 1.0, 12, x=8.0, w=1.2)

        issues = _analyze_text_shapes(presentation)

        self.assertNotIn("title_clearance_risk", {item["type"] for item in issues})


if __name__ == "__main__":
    unittest.main()
