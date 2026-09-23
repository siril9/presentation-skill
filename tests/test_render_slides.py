from __future__ import annotations

import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import render_cache
import render_slides


class RenderCacheTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.input = self.root / "deck.pptx"
        self.output = self.root / "renders"
        self.cache = self.root / "cache"
        self.version = "1.0"
        self.tool_path = "/tools/"
        self.slide_count = 2
        self.write_deck("first")
        self.addCleanup(patch.stopall)
        patch.object(render_slides.shutil, "which", side_effect=lambda name: None if name == "unoconvert" else f"/tools/{name}").start()
        patch.object(render_slides, "_tool_identity", side_effect=lambda name, flag: {
            "path": self.tool_path + name, "version": self.version, "sha256": "tool-bytes",
        }).start()
        self.fresh = patch.object(render_slides, "_fresh_render", side_effect=self.fake_render).start()

    def write_deck(self, value: str, external: bool = False) -> None:
        with zipfile.ZipFile(self.input, "w") as package:
            package.writestr("ppt/presentation.xml", f"<presentation>{value}</presentation>")
            if external:
                package.writestr("ppt/_rels/presentation.xml.rels", '<Relationships><Relationship TargetMode="External" Type="image" Target="file.png"/></Relationships>')

    def fake_render(self, source: Path, directory: Path, dpi: int, image_format: str, **kwargs) -> list[Path]:
        paths = []
        for index in range(1, self.slide_count + 1):
            path = directory / f"slide-{index:02d}{'.png' if image_format == 'png' else '.jpg'}"
            Image.new("RGB", (16, 9), (index, dpi % 256, 0)).save(path)
            paths.append(path)
        return paths

    def run_render(self, *extra: str, use_cache: bool = True) -> dict:
        argv = ["render_slides.py", "--input", str(self.input), "--outdir", str(self.output)]
        if use_cache:
            argv += ["--cache-dir", str(self.cache)]
        argv += list(extra)
        with patch.object(sys, "argv", argv), contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(render_slides.main(), 0)
        report_path = Path(extra[extra.index("--report") + 1]) if "--report" in extra else self.output / "render_report.json"
        return json.loads(report_path.read_text())

    def entry(self, report: dict) -> Path:
        return self.cache / report["cache"]["key"]

    def test_exact_hit_survives_output_cleanup_and_input_rename(self) -> None:
        cold = self.run_render()
        shutil.rmtree(self.output)
        renamed = self.root / "renamed.pptx"
        self.input.rename(renamed)
        self.input = renamed
        warm = self.run_render()
        self.assertEqual(cold["cache"]["status"], "miss")
        self.assertEqual(warm["cache"]["status"], "hit")
        self.assertEqual(self.fresh.call_count, 1)
        self.assertEqual(cold["images"], warm["images"])

    def test_content_change_with_same_size_and_mtime_misses(self) -> None:
        cold = self.run_render()
        stat = self.input.stat()
        self.write_deck("other")
        self.assertEqual(self.input.stat().st_size, stat.st_size)
        os.utime(self.input, ns=(stat.st_atime_ns, stat.st_mtime_ns))
        changed = self.run_render()
        self.assertEqual(changed["cache"]["status"], "miss")
        self.assertNotEqual(cold["cache"]["key"], changed["cache"]["key"])

    def test_format_and_dpi_invalidate(self) -> None:
        reports = [self.run_render(), self.run_render("--dpi", "180"), self.run_render("--format", "png")]
        self.assertEqual(len({r["cache"]["key"] for r in reports}), 3)
        self.assertEqual(self.fresh.call_count, 3)
        self.assertFalse(list(self.output.glob("*.jpg")))

    def test_converter_version_and_identity_invalidate(self) -> None:
        reports = [self.run_render()]
        self.version = "2.0"
        reports.append(self.run_render())
        self.tool_path = "/replacement/"
        reports.append(self.run_render())
        self.assertEqual(len({r["cache"]["key"] for r in reports}), 3)
        self.assertEqual(self.fresh.call_count, 3)

    def test_helper_code_invalidation(self) -> None:
        old = self.run_render()
        original = render_slides.file_hash
        for helper in ("render_slides.py", "render_cache.py"):
            with self.subTest(helper=helper), patch.object(render_slides, "file_hash", side_effect=lambda p: "changed-helper" if p.name == helper else original(p)):
                new = self.run_render()
                self.assertNotEqual(old["cache"]["key"], new["cache"]["key"])
                self.assertEqual(new["cache"]["status"], "miss")

    def test_environment_change_invalidates(self) -> None:
        original = self.run_render()
        with patch.dict(os.environ, {"SAL_USE_VCLPLUGIN": "changed-backend"}):
            changed = self.run_render()
        self.assertNotEqual(original["cache"]["key"], changed["cache"]["key"])
        self.assertEqual(changed["cache"]["status"], "miss")

    def test_custom_report_and_visual_prompt_on_hit(self) -> None:
        self.run_render()
        report = self.root / "custom.json"
        self.run_render("--report", str(report), "--emit-visual-prompt")
        payload = json.loads(report.read_text())
        self.assertEqual(payload["schema_version"], "slide_render_report_v1")
        self.assertEqual(payload["cache"]["status"], "hit")
        self.assertEqual(len(payload["images"]), payload["page_count"])
        for image in payload["images"]:
            self.assertEqual(render_cache.file_hash(Path(image["path"])), image["sha256"])

    def test_corrupt_same_size_image_is_not_reused(self) -> None:
        entry = self.entry(self.run_render())
        image = entry / "slide-01.jpg"
        stat = image.stat()
        data = bytearray(image.read_bytes())
        data[-1] ^= 1
        image.write_bytes(data)
        os.utime(image, ns=(stat.st_atime_ns, stat.st_mtime_ns))
        self.assertEqual(self.run_render()["cache"]["status"], "miss")
        self.assertEqual(self.run_render()["cache"]["status"], "hit")
        self.assertEqual(self.fresh.call_count, 2)

    def test_missing_manifest_missing_image_and_extra_image_miss(self) -> None:
        for defect in ("manifest", "image", "extra", "malformed", "page_count", "traversal", "symlink"):
            with self.subTest(defect=defect):
                entry = self.entry(self.run_render())
                manifest = entry / "manifest.json"
                if defect == "manifest":
                    manifest.unlink()
                elif defect == "image":
                    (entry / "slide-02.jpg").unlink()
                elif defect == "extra":
                    (entry / "slide-03.jpg").write_bytes(b"stale")
                elif defect == "malformed":
                    manifest.write_text("{")
                elif defect == "symlink":
                    image = entry / "slide-01.jpg"
                    image.unlink()
                    image.symlink_to(self.output / image.name)
                else:
                    payload = json.loads(manifest.read_text())
                    if defect == "page_count":
                        payload["page_count"] += 1
                    else:
                        payload["files"][0]["name"] = "../outside.jpg"
                    manifest.write_text(json.dumps(payload))
                self.assertEqual(self.run_render()["cache"]["status"], "miss")

    def test_no_cache_bypasses_reads_writes_and_identity_probes(self) -> None:
        self.run_render()
        with patch.object(render_slides, "_cache_inputs", side_effect=AssertionError("cache probe")), patch.object(render_slides, "publish", side_effect=AssertionError("cache write")):
            report = self.run_render("--no-cache")
        self.assertEqual(report["cache"]["reason"], "no_cache")
        self.assertEqual(self.fresh.call_count, 2)
        self.assertEqual(self.run_render()["cache"]["status"], "hit")

    def test_default_is_opt_in(self) -> None:
        report = self.run_render(use_cache=False)
        self.assertEqual(report["cache"]["status"], "disabled")
        self.assertFalse(self.cache.exists())

    def test_unverifiable_identity_and_external_resources_bypass(self) -> None:
        with patch.object(render_slides.shutil, "which", return_value="/tools/binary"):
            report = self.run_render()
            self.assertEqual(report["cache"]["status"], "bypassed")
            self.assertIn("daemon", report["cache"]["reason"])
        self.write_deck("first", external=True)
        self.assertIn("external", self.run_render()["cache"]["reason"])
        self.assertFalse(self.cache.exists())

    def test_unknown_tool_version_bypasses(self) -> None:
        with patch.object(render_slides, "_tool_identity", side_effect=RuntimeError("unknown version")):
            self.assertEqual(self.run_render()["cache"]["status"], "bypassed")

    def test_failed_conversion_does_not_publish_or_leave_success_report(self) -> None:
        report = self.run_render()
        self.write_deck("other")
        with patch.object(render_slides, "_fresh_render", side_effect=RuntimeError("conversion failed")):
            with self.assertRaisesRegex(RuntimeError, "conversion failed"):
                self.run_render()
        self.assertFalse((self.output / "render_report.json").exists())
        self.assertEqual(len(list(self.cache.glob("*/manifest.json"))), 1)
        self.assertTrue(self.entry(report).is_dir())

    def test_input_mutation_during_conversion_does_not_publish(self) -> None:
        def mutate(*args, **kwargs):
            paths = self.fake_render(*args, **kwargs)
            self.write_deck("changed")
            return paths
        self.fresh.side_effect = mutate
        with self.assertRaisesRegex(RuntimeError, "Input changed"):
            self.run_render()
        self.assertFalse(list(self.cache.glob("*/manifest.json")))

    def test_dependency_change_during_conversion_does_not_publish(self) -> None:
        def mutate(*args, **kwargs):
            paths = self.fake_render(*args, **kwargs)
            self.version = "changed"
            return paths
        self.fresh.side_effect = mutate
        with self.assertRaisesRegex(RuntimeError, "dependencies changed"):
            self.run_render()
        self.assertFalse(list(self.cache.glob("*/manifest.json")))

    def test_output_cleanup_removes_trailing_and_other_format_slides_only(self) -> None:
        self.run_render()
        (self.output / "slide-10.png").write_bytes(b"old")
        (self.output / "notes.txt").write_text("keep")
        self.slide_count = 1
        self.write_deck("short")
        report = self.run_render()
        self.assertEqual(report["page_count"], 1)
        self.assertEqual(sorted(p.name for p in self.output.glob("slide-*")), ["slide-01.jpg"])
        self.assertEqual((self.output / "notes.txt").read_text(), "keep")

    def test_nested_cache_is_rejected(self) -> None:
        self.cache = self.output / "cache"
        with self.assertRaisesRegex(ValueError, "non-nested"):
            self.run_render()

    def test_interrupted_cache_publish_is_not_visible(self) -> None:
        original = render_cache.os.replace
        def interrupt(source, target):
            if Path(source).name == "entry":
                raise OSError("interrupted publication")
            return original(source, target)
        with patch.object(render_cache.os, "replace", side_effect=interrupt):
            report = self.run_render()
        self.assertIn("cache_write_failed", report["cache"]["reason"])
        self.assertFalse(list(self.cache.glob("*/manifest.json")))
        self.assertEqual(self.run_render()["cache"]["status"], "miss")


class RenderValidationTests(unittest.TestCase):
    def test_tool_identity_hashes_bytes_even_when_version_unchanged(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            binary = Path(temporary) / "pdftoppm"
            binary.write_bytes(b"first")
            response = subprocess.CompletedProcess([], 0, "", "version 1.0")
            with patch.object(render_slides.shutil, "which", return_value=str(binary)), patch.object(render_slides.subprocess, "run", return_value=response):
                first = render_slides._tool_identity("pdftoppm", "-v")
                binary.write_bytes(b"other")
                second = render_slides._tool_identity("pdftoppm", "-v")
                self.assertNotEqual(first["sha256"], second["sha256"])
                self.assertEqual(first["version"], second["version"])
                response.returncode = 1
                with self.assertRaisesRegex(RuntimeError, "version"):
                    render_slides._tool_identity("pdftoppm", "-v")

    def test_hash_rejects_path_replaced_while_reading(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "input"
            replacement = Path(temporary) / "replacement"
            path.write_bytes(b"same")
            replacement.write_bytes(b"same")
            with patch.object(Path, "stat", return_value=replacement.stat()):
                with self.assertRaisesRegex(RuntimeError, "changed while hashing"):
                    render_cache.file_hash(path)

    def test_isolated_converter_uses_fresh_profile_and_skips_daemon(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            def convert(command):
                self.assertIn(f"-env:UserInstallation={(root / 'lo-profile').resolve().as_uri()}", command)
                (root / "deck.pdf").write_bytes(b"pdf")
            with patch.object(render_slides.subprocess, "run", side_effect=AssertionError("daemon")), patch.object(render_slides, "_run", side_effect=convert):
                render_slides._render_with_daemon_or_fallback(root / "deck.pptx", root, isolated=True)

    def test_uncached_fallback_also_isolates_office_profile(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            def convert(command):
                self.assertIn(f"-env:UserInstallation={(root / 'lo-profile').resolve().as_uri()}", command)
                (root / "deck.pdf").write_bytes(b"pdf")
            with patch.object(render_slides.shutil, "which", return_value=None), patch.object(render_slides, "_run", side_effect=convert):
                render_slides._render_with_daemon_or_fallback(root / "deck.pptx", root, isolated=False)

    def test_partial_or_corrupt_images_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            image = root / "slide-1.png"
            Image.new("RGB", (20, 10)).save(image)
            response = subprocess.CompletedProcess([], 0, "Pages: 2\n", "")
            with patch.object(render_slides.subprocess, "run", return_value=response):
                with self.assertRaisesRegex(RuntimeError, "image count"):
                    render_slides._validate_images(root / "deck.pdf", [image], "png")
                response.stdout = "Pages: 1\n"
                render_slides._validate_images(root / "deck.pdf", [image], "png")
                image.write_bytes(b"truncated")
                with self.assertRaises(OSError):
                    render_slides._validate_images(root / "deck.pdf", [image], "png")

    def test_valid_but_short_pdf_is_rejected_against_pptx(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            presentation = Presentation()
            for _ in range(2):
                presentation.slides.add_slide(presentation.slide_layouts[6])
            pptx = root / "deck.pptx"
            presentation.save(pptx)
            Image.new("RGB", (20, 10)).save(root / "slide-1.png")
            with patch.object(render_slides, "_render_with_daemon_or_fallback", return_value=root / "deck.pdf"), patch.object(render_slides, "_run"):
                with self.assertRaisesRegex(RuntimeError, "visible PPTX slides"):
                    render_slides._fresh_render(pptx, root, 90, "png")

    def test_failed_daemon_pdf_removed_before_fallback(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            pdf = root / "deck.pdf"
            def daemon(*args, **kwargs):
                pdf.write_bytes(b"partial")
                return subprocess.CompletedProcess([], 1, "", "failed")
            def fallback(command):
                self.assertFalse(pdf.exists())
                pdf.write_bytes(b"fresh")
            with patch.object(render_slides.shutil, "which", return_value="/tool"), patch.object(render_slides.subprocess, "run", side_effect=daemon), patch.object(render_slides, "_run", side_effect=fallback):
                self.assertEqual(render_slides._render_with_daemon_or_fallback(root / "deck.pptx", root).read_bytes(), b"fresh")


if __name__ == "__main__":
    unittest.main()
